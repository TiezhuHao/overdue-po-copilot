from collections import Counter
from datetime import date, datetime, timezone
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import pytest
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.reporting.excel_export import ReportExcelExporter
from app.reporting.report_header_manifest import get_report_manifest
from app.models.platform import DatasetVersion
from app.services.report_queries import REPORTS
from tests.test_forecast_integration import forecast_worlds
from tests.test_operational_integration import operational_worlds


pytestmark = pytest.mark.integration


def _read_export(path, manifest):
    header_end = max(manifest["sheet_structure"][0]["header_rows"])
    count = len(manifest["columns"])
    with path.open("rb") as source:
        workbook = load_workbook(source, data_only=False)
        try:
            sheets = []
            for sheet in workbook.worksheets:
                headers = [
                    sheet[column["slot"]["header_cell"] if column["slot"]
                          else column["source_header_cells"][-1]].value
                    for column in manifest["columns"]
                ]
                rows = [list(row) for row in sheet.iter_rows(
                    min_row=header_end + 1, max_row=sheet.max_row, max_col=count, values_only=True,
                )] if sheet.max_row > header_end else []
                sheets.append({
                    "name": sheet.title, "headers": headers, "rows": rows,
                    "column_count": sheet.max_column, "row_count": len(rows),
                    "merges": sorted(str(value) for value in sheet.merged_cells.ranges),
                    "freeze_panes": sheet.freeze_panes, "auto_filter": sheet.auto_filter.ref,
                    "widths": [sheet.column_dimensions[get_column_letter(i)].width for i in range(1, count + 1)],
                    "header_bold": all(
                        sheet[column["slot"]["header_cell"] if column["slot"]
                              else column["source_header_cells"][-1]].font.bold
                        for column in manifest["columns"]
                    ),
                    "data_formats": [
                        sheet.cell(header_end + 1, i).number_format for i in range(1, count + 1)
                    ] if rows else [],
                    "formula_count": sum(cell.data_type == "f" for row in sheet for cell in row),
                    "error_count": sum(cell.data_type == "e" for row in sheet for cell in row),
                })
            return sheets
        finally:
            workbook.close()


@pytest.fixture(scope="module")
def excel_exports(migrated_database, operational_worlds, tmp_path_factory):
    world = operational_worlds[0]
    output = tmp_path_factory.mktemp("six-report-exports")
    results, inspections, prepared, expected_counts, repeated = {}, {}, {}, {}, {}
    with pytest.MonkeyPatch.context() as patch:
        # A working Python process needs neither Node nor inherited runtime settings.
        patch.setenv("PATH", "")
        patch.delenv("REPORT_EXPORT_NODE", raising=False)
        patch.delenv("REPORT_EXPORT_NODE_MODULES", raising=False)
        with Session(migrated_database) as session:
            exporter = ReportExcelExporter(session)
            dataset_name = session.get(DatasetVersion, world.dataset_version_id).version_name
            for report_id in range(1, 7):
                prepared[report_id] = exporter.prepare(report_id, dataset_name)
                raw_count = session.scalar(text(
                    f"SELECT count(*) FROM reporting.{REPORTS[report_id]['view']} WHERE dataset_version_id=:did"
                ), {"did": world.dataset_version_id})
                expected_counts[report_id] = raw_count // 7 if report_id == 3 else raw_count
                if report_id == 3:
                    assert raw_count % 7 == 0
                results[report_id] = exporter.export(report_id, dataset_name, output)
                manifest = get_report_manifest(report_id)
                inspections[report_id] = _read_export(results[report_id].output_path, manifest)
                second = exporter.export(report_id, dataset_name, output / "repeat")
                repeated[report_id] = _read_export(second.output_path, manifest)
    return results, inspections, prepared, expected_counts, repeated


@pytest.mark.parametrize("report_id,count", [(1, 37), (2, 106), (3, 23), (4, 44), (5, 22), (6, 60)])
def test_excel_headers_exact(report_id, count, excel_exports):
    _, inspections, prepared, _, _ = excel_exports
    manifest = get_report_manifest(report_id)
    assert len(manifest["columns"]) == count
    for sheet, source in zip(inspections[report_id], prepared[report_id]["sheets"], strict=True):
        expected = [source["header_overrides"].get(column["canonical_field"], column["display_name"])
                    for column in manifest["columns"]]
        assert sheet["headers"] == expected
        assert sheet["column_count"] == count
        assert sheet["merges"] == sorted(manifest["sheet_structure"][0]["merges"])


@pytest.mark.parametrize("report_id", range(1, 7))
def test_excel_row_counts_and_reopen(report_id, excel_exports):
    results, inspections, _, expected_counts, _ = excel_exports
    result = results[report_id]
    assert result.output_path.is_file() and result.output_path.stat().st_size > 0
    assert result.row_count == sum(sheet["row_count"] for sheet in inspections[report_id]) == expected_counts[report_id]
    assert len(inspections[report_id]) == result.sheet_count


def test_report3_excel_pivot_and_dynamic_months(excel_exports):
    results, inspections, prepared, _, _ = excel_exports
    assert results[3].sheet_count == len(prepared[3]["sheets"]) > 1
    for sheet in inspections[3]:
        headers = sheet["headers"][14:21]
        assert all(re.fullmatch(r"\d{6}", value) for value in headers)
        month_indexes = [int(value[:4]) * 12 + int(value[4:]) for value in headers]
        assert month_indexes == list(range(month_indexes[0], month_indexes[0] + 7))


def test_report3_pivot_totals_match_canonical_long(migrated_database, operational_worlds, excel_exports):
    with Session(migrated_database) as session:
        expected = list(session.execute(text(
            "SELECT forecast_version_name,forecast_month,sum(forecast_qty) quantity "
            "FROM reporting.report3_forecast_history WHERE dataset_version_id=:did "
            "GROUP BY forecast_version_name,forecast_month"
        ), {"did": operational_worlds[0].dataset_version_id}).mappings())
    totals = Counter()
    columns = get_report_manifest(3)["columns"]
    version_index = next(i for i, column in enumerate(columns) if column["canonical_field"] == "forecast_version_name")
    for sheet in excel_exports[1][3]:
        for row in sheet["rows"]:
            for index in range(14, 21):
                totals[(row[version_index], sheet["headers"][index])] += row[index]
    assert len(totals) == len(expected)
    for row in expected:
        assert totals[(row["forecast_version_name"], row["forecast_month"].strftime("%Y%m"))] == pytest.approx(float(row["quantity"]))


def test_report4_excel_thirteen_monday_weeks(excel_exports):
    headers = excel_exports[1][4][0]["headers"][29:42]
    days = [date.fromisoformat(value) for value in headers]
    assert len(days) == 13 and all(day.weekday() == 0 for day in days)
    assert all((later - earlier).days == 7 for earlier, later in zip(days, days[1:]))


def test_report6_excel_six_natural_months(excel_exports):
    headers = excel_exports[1][6][0]["headers"][20:26]
    assert headers == ["202609", "202610", "202611", "202612", "202701", "202702"]


def test_excel_unknown_business_fields_remain_blank(excel_exports):
    _, inspections, _, _, _ = excel_exports
    for report_id in (2, 4, 6):
        manifest = get_report_manifest(report_id)
        deferred = [index for index, column in enumerate(manifest["columns"])
                    if column["business_status"] == "DEFERRED_BUSINESS_FORMULA"
                    or column["canonical_field"] == "kd"]
        assert deferred
        for sheet in inspections[report_id]:
            for row in sheet["rows"]:
                assert all(row[index] is None for index in deferred)


def _comparable(value):
    if isinstance(value, str) and re.match(r"^\d{4}-\d{2}-\d{2}T", value):
        value = datetime.fromisoformat(value)
    if isinstance(value, datetime):
        if value.tzinfo is not None:
            value = value.astimezone(timezone.utc).replace(tzinfo=None)
        return value.date().isoformat() if value.time().isoformat() == "00:00:00" else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, list):
        return ",".join("" if item is None else str(item) for item in value) or None
    return None if value == "" else value


@pytest.mark.parametrize("report_id", range(1, 7))
def test_excel_roundtrip_all_values(report_id, excel_exports):
    _, inspections, prepared, _, _ = excel_exports
    for sheet, source in zip(inspections[report_id], prepared[report_id]["sheets"], strict=True):
        for actual, expected in zip(sheet["rows"], source["rows"], strict=True):
            for actual_cell, expected_cell in zip(actual, expected, strict=True):
                if isinstance(expected_cell, (float, int)) and not isinstance(expected_cell, bool):
                    assert actual_cell == pytest.approx(expected_cell, rel=1e-12, abs=1e-8)
                else:
                    assert _comparable(actual_cell) == _comparable(expected_cell)


@pytest.mark.parametrize("report_id", range(1, 7))
def test_excel_business_content_and_order_are_deterministic(report_id, excel_exports):
    assert excel_exports[1][report_id] == excel_exports[4][report_id]


@pytest.mark.parametrize("report_id", range(1, 7))
def test_excel_basic_formatting(report_id, excel_exports):
    manifest = get_report_manifest(report_id)
    header_end = max(manifest["sheet_structure"][0]["header_rows"])
    for sheet in excel_exports[1][report_id]:
        assert sheet["header_bold"]
        assert sheet["freeze_panes"] == f"A{header_end + 1}"
        assert sheet["auto_filter"].startswith(f"A{header_end}:")
        assert all(10 <= width <= 28 for width in sheet["widths"])
        for i, column in enumerate(manifest["columns"]):
            if not sheet["rows"]:
                continue
            field = column["canonical_field"]
            if field.endswith(("date", "_at")):
                assert sheet["data_formats"][i] == "yyyy-mm-dd"
            elif "qty" in field:
                assert sheet["data_formats"][i] == "#,##0.0000"


def test_excel_no_formulas_errors_or_proprietary_terms(excel_exports):
    term_file = Path(__file__).resolve().parents[2] / "scripts" / "forbidden_terms.txt"
    terms = [line.strip().casefold() for line in term_file.read_text(encoding="utf-8").splitlines()
             if line.strip() and not line.lstrip().startswith("#") and line.strip() not in {"BG", "PDT", "PCBA"}]
    errors = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!"}
    for sheets in excel_exports[1].values():
        for sheet in sheets:
            assert sheet["formula_count"] == sheet["error_count"] == 0
            for value in sheet["headers"] + [cell for row in sheet["rows"] for cell in row]:
                if isinstance(value, str):
                    assert value not in errors
                    assert not any(term in value.casefold() for term in terms)
