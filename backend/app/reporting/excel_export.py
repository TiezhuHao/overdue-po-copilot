from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import text

from app.reporting.report_header_manifest import get_report_manifest
from app.services.report_queries import REPORTS, ReportQueryService


REPORT_FILENAMES = {
    1: "report1_overdue_po_detail.xlsx",
    2: "report2_material_supply_demand.xlsx",
    3: "report3_forecast_history.xlsx",
    4: "report4_latest_13w_forecast.xlsx",
    5: "report5_product_configuration.xlsx",
    6: "report6_stockpile_detail.xlsx",
}


class ReportExcelExportError(RuntimeError):
    pass


@dataclass(frozen=True)
class ExportResult:
    report_id: int
    output_path: Path
    row_count: int
    sheet_count: int


def _json_value(value):
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, UUID):
        return str(value)
    if isinstance(value, tuple):
        return [_json_value(item) for item in value]
    if isinstance(value, list):
        return [_json_value(item) for item in value]
    return value


def _month_add(day, offset):
    month = day.month - 1 + offset
    return date(day.year + month // 12, month % 12 + 1, 1)


class ReportExcelExporter:
    def __init__(self, executor):
        self.executor = executor

    def _dataset(self, name):
        return ReportQueryService(self.executor).select_dataset(None, name)

    def _static_rows(self, report_id, dataset_id):
        definition = REPORTS[report_id]
        return [dict(row) for row in self.executor.execute(text(
            f"SELECT * FROM reporting.{definition['view']} WHERE dataset_version_id=:did ORDER BY {definition['order']}"
        ), {"did": dataset_id}).mappings()]

    def _report3_sheets(self, dataset):
        rows = self._static_rows(3, dataset.dataset_version_id)
        fields = [column["canonical_field"] for column in get_report_manifest(3)["columns"]]
        static = [field for field in fields if not field.startswith("forecast_qty_m")]
        grouped = {}
        for row in rows:
            key = (row["po_line_schedule_id"], row["material_id"], row["project_id"], row["forecast_version_id"])
            item = grouped.setdefault(key, {field: row.get(field) for field in static})
            index = (row["forecast_month"].year - row["forecast_version_date"].year) * 12 + row["forecast_month"].month - row["forecast_version_date"].month
            if 0 <= index <= 6:
                item[f"forecast_qty_m{index}"] = row["forecast_qty"]
        by_version = defaultdict(list)
        for item in grouped.values():
            by_version[(item["forecast_version_date"], item["forecast_version_name"])].append(item)
        sheets = []
        for (version_date, version_name), values in sorted(by_version.items()):
            overrides = {f"forecast_qty_m{i}": _month_add(version_date.replace(day=1), i).strftime("%Y%m") for i in range(7)}
            safe_name = f"{version_date:%Y%m%d}-{version_name}"[:31].replace("/", "-").replace("\\", "-")
            sheets.append(self._sheet_payload(3, safe_name, values, overrides))
        return sheets

    def _report6_rows(self, dataset):
        rows = self._static_rows(6, dataset.dataset_version_id)
        forecasts = defaultdict(dict)
        for row in self.executor.execute(text(
            "SELECT material_id,month_index,forecast_month,forecast_qty FROM reporting.report6_stockpile_forecast_long "
            "WHERE dataset_version_id=:did ORDER BY material_id,month_index"
        ), {"did": dataset.dataset_version_id}).mappings():
            forecasts[row["material_id"]][row["month_index"]] = row
        ages = defaultdict(dict)
        for row in self.executor.execute(text(
            "SELECT material_id,age_threshold_days,age_qty FROM reporting.report6_stockpile_age_long "
            "WHERE dataset_version_id=:did ORDER BY material_id,age_threshold_days"
        ), {"did": dataset.dataset_version_id}).mappings():
            ages[row["material_id"]][row["age_threshold_days"]] = row["age_qty"]
        for row in rows:
            for index, item in forecasts[row["material_id"]].items():
                row[f"future_month_{index:02d}_demand_qty"] = item["forecast_qty"]
            for threshold in (30, 60, 90, 120, 150, 180, 270, 365):
                row[f"inventory_age_over_{threshold}d_qty"] = ages[row["material_id"]].get(threshold)
        return rows

    def _sheet_payload(self, report_id, sheet_name, rows, header_overrides=None):
        manifest = get_report_manifest(report_id)
        columns = manifest["columns"]
        return {
            "name": sheet_name,
            "fields": [column["canonical_field"] for column in columns],
            "rows": [[_json_value(row.get(column["canonical_field"])) for column in columns] for row in rows],
            "header_overrides": header_overrides or {},
        }

    def prepare(self, report_id, dataset_version_name):
        dataset = self._dataset(dataset_version_name)
        manifest = get_report_manifest(report_id)
        if report_id == 3:
            sheets = self._report3_sheets(dataset)
        elif report_id == 6:
            rows = self._report6_rows(dataset)
            version_date = rows[0]["stockpile_version_date"] if rows else dataset.snapshot_date
            overrides = {f"future_month_{i:02d}_demand_qty": _month_add(version_date.replace(day=1), i).strftime("%Y%m") for i in range(1, 7)}
            sheets = [self._sheet_payload(6, manifest["sheet_name"], rows, overrides)]
        else:
            rows = self._static_rows(report_id, dataset.dataset_version_id)
            overrides = {}
            if report_id == 4:
                start = dataset.snapshot_date + timedelta(days=(7 - dataset.snapshot_date.weekday()) % 7)
                overrides = {f"week_{i:02d}_forecast_qty": (start + timedelta(days=7 * (i - 1))).isoformat() for i in range(1, 14)}
            sheets = [self._sheet_payload(report_id, manifest["sheet_name"], rows, overrides)]
        return {"report_id": report_id, "manifest": manifest, "sheets": sheets}

    @staticmethod
    def _display_value(field, value):
        if isinstance(value, list):
            # Preserve the previous comma-separated display of canonical arrays.
            return ",".join("" if item is None else str(item) for item in value)
        if isinstance(value, str) and field.endswith(("date", "_at")):
            try:
                if len(value) == 10:
                    return date.fromisoformat(value)
                timestamp = datetime.fromisoformat(value)
                # Excel has no timezone type; preserve the old writer's UTC value.
                if timestamp.tzinfo is not None:
                    timestamp = timestamp.astimezone(timezone.utc).replace(tzinfo=None)
                return timestamp
            except ValueError:
                pass
        return value

    @staticmethod
    def _set_value(cell, value):
        cell.value = value
        if isinstance(value, str):
            # Report text must remain literal, including formula-like source text.
            cell.data_type = "s"

    def _write_sheet(self, workbook, manifest, data_sheet):
        sheet = workbook.create_sheet(data_sheet["name"])
        sheet.sheet_view.showGridLines = False
        structure = manifest["sheet_structure"][0]
        header_end = max(structure["header_rows"])
        columns = manifest["columns"]
        last_column = get_column_letter(len(columns))
        for address, value in structure["header_cells"].items():
            self._set_value(sheet[address], value)
        for merge in structure["merges"]:
            sheet.merge_cells(merge)
        for column in columns:
            field = column["canonical_field"]
            if column["slot"] and field in data_sheet["header_overrides"]:
                self._set_value(sheet[column["slot"]["header_cell"]], data_sheet["header_overrides"][field])

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="1F4E78")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin = Side(style="thin", color="B4C7E7")
        for row in sheet.iter_rows(min_row=1, max_row=header_end, max_col=len(columns)):
            sheet.row_dimensions[row[0].row].height = 30
            for cell in row:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                if cell.row == header_end:
                    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for index, column in enumerate(columns, 1):
            field = column["canonical_field"]
            if field.endswith(("date", "_at")):
                number_format = "yyyy-mm-dd"
            elif "ratio" in field or "completion" in field:
                number_format = "0.00%"
            elif any(term in field for term in ("qty", "amount", "price")):
                number_format = "#,##0.0000"
            else:
                number_format = "General"
            for row_number, values in enumerate(data_sheet["rows"], header_end + 1):
                cell = sheet.cell(row_number, index)
                self._set_value(cell, self._display_value(field, values[index - 1]))
                cell.number_format = number_format
            sample = [str(self._display_value(field, row[index - 1]) or "")
                      for row in data_sheet["rows"][:50]]
            width = max(10, len(" ".join(column["header_path"])) + 2,
                        *(len(value) + 2 for value in sample))
            sheet.column_dimensions[get_column_letter(index)].width = min(28, width)

        sheet.freeze_panes = f"A{header_end + 1}"
        last_row = max(header_end + 1, header_end + len(data_sheet["rows"]))
        sheet.auto_filter.ref = f"A{header_end}:{last_column}{last_row}"

    def export(self, report_id, dataset_version_name, output_dir):
        if report_id not in REPORT_FILENAMES:
            raise ReportExcelExportError("UNKNOWN_REPORT")
        payload = self.prepare(report_id, dataset_version_name)
        output_dir = Path(output_dir).resolve()
        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = output_dir / REPORT_FILENAMES[report_id]
        workbook = Workbook()
        workbook.remove(workbook.active)
        workbook.properties.creator = "System A"
        workbook.properties.lastModifiedBy = "System A"
        workbook.properties.description = "Synthetic data only."
        try:
            for data_sheet in payload["sheets"]:
                self._write_sheet(workbook, payload["manifest"], data_sheet)
            workbook.save(output_path)
        finally:
            workbook.close()
        return ExportResult(
            report_id, output_path,
            sum(len(sheet["rows"]) for sheet in payload["sheets"]),
            len(payload["sheets"]),
        )
