from datetime import datetime

from openpyxl import load_workbook
import pytest

from app.reporting.excel_export import ReportExcelExporter
from app.reporting.report_header_manifest import get_report_manifest


@pytest.mark.parametrize("source_text", ["=1+1", "+SUM(A1:A2)", "@SUM(A1:A2)", "#REF!"])
def test_source_text_remains_literal_in_workbook(tmp_path, monkeypatch, source_text):
    manifest = get_report_manifest(1)
    fields = [column["canonical_field"] for column in manifest["columns"]]
    index = fields.index("material_description")
    row = [None] * len(fields)
    row[index] = source_text
    exporter = ReportExcelExporter(None)
    monkeypatch.setattr(exporter, "prepare", lambda *_: {
        "manifest": manifest,
        "sheets": [{"name": manifest["sheet_name"], "fields": fields, "rows": [row], "header_overrides": {}}],
    })
    result = exporter.export(1, "synthetic-test", tmp_path)
    workbook = load_workbook(result.output_path, data_only=False)
    try:
        cell = workbook.worksheets[0].cell(2, index + 1)
        assert cell.value == source_text
        assert cell.data_type == "s"
        assert workbook.properties.creator == "System A"
    finally:
        workbook.close()


@pytest.mark.parametrize("source,expected", [
    ("2025-07-03", datetime(2025, 7, 3)),
    ("2025-07-03T12:00:00+00:00", datetime(2025, 7, 3, 12)),
    ("2025-07-03T20:00:00+08:00", datetime(2025, 7, 3, 12)),
])
def test_dates_and_timestamps_remain_native_excel_values(tmp_path, monkeypatch, source, expected):
    manifest = get_report_manifest(1)
    fields = [column["canonical_field"] for column in manifest["columns"]]
    index = fields.index("completion_at")
    row = [None] * len(fields)
    row[index] = source
    exporter = ReportExcelExporter(None)
    monkeypatch.setattr(exporter, "prepare", lambda *_: {
        "manifest": manifest,
        "sheets": [{"name": manifest["sheet_name"], "fields": fields, "rows": [row], "header_overrides": {}}],
    })
    result = exporter.export(1, "synthetic-test", tmp_path)
    workbook = load_workbook(result.output_path)
    try:
        cell = workbook.worksheets[0].cell(2, index + 1)
        assert cell.value == expected
        assert cell.data_type == "d"
        assert cell.number_format == "yyyy-mm-dd"
    finally:
        workbook.close()
